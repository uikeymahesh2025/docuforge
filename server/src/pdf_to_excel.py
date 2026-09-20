import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pymupdf as fitz

def convert_pdf_to_excel(pdf_path: str, excel_path: str):
    pdf_doc = fitz.open(pdf_path)
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    table_count = 0

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid") # Dark Navy
    cell_font = Font(name="Calibri", size=10)
    thin_border_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for page_idx, page in enumerate(pdf_doc):
        # 1. Try PyMuPDF native table detector
        tabs = []
        try:
            tabs = page.find_tables().tables
        except Exception:
            tabs = []

        if tabs:
            for t_idx, t in enumerate(tabs):
                extracted = t.extract()
                if not extracted or len(extracted) == 0:
                    continue
                
                table_count += 1
                sheet_title = f"P{page_idx + 1}_Table{t_idx + 1}"
                ws = wb.create_sheet(title=sheet_title[:31])

                for r_idx, row in enumerate(extracted):
                    for c_idx, val in enumerate(row):
                        cell_val = "" if val is None else str(val).strip()
                        # Clean multiple newlines in cell
                        cell_val = " ".join(cell_val.split())
                        
                        # Try to cast numeric values
                        numeric_val = None
                        try:
                            clean_num = cell_val.replace(",", "").replace("$", "").replace("₹", "")
                            if clean_num and (clean_num.replace(".", "", 1).isdigit() or (clean_num.startswith("-") and clean_num[1:].replace(".", "", 1).isdigit())):
                                numeric_val = float(clean_num) if "." in clean_num else int(clean_num)
                        except Exception:
                            numeric_val = None

                        cell = ws.cell(row=r_idx + 1, column=c_idx + 1, value=numeric_val if numeric_val is not None else cell_val)
                        cell.border = cell_border

                        if r_idx == 0:
                            cell.font = header_font
                            cell.fill = header_fill
                            cell.alignment = Alignment(horizontal="center" if numeric_val is not None else "left", vertical="center")
                        else:
                            cell.font = cell_font
                            cell.alignment = Alignment(horizontal="right" if numeric_val is not None else "left", vertical="center")
                            if r_idx % 2 == 0:
                                cell.fill = alt_fill

                # Auto-adjust column widths
                for col in ws.columns:
                    col_letter = get_column_letter(col[0].column)
                    max_len = max(len(str(c.value or "")) for c in col)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        else:
            # 2. Heuristic text-block table extractor for borderless tables
            text_dict = page.get_text("dict")
            raw_spans = []
            for b in text_dict.get("blocks", []):
                if b.get("type") == 0:
                    for l in b.get("lines", []):
                        for s in l.get("spans", []):
                            t = s.get("text", "").strip()
                            if t:
                                raw_spans.append({
                                    "text": t,
                                    "x0": s["bbox"][0],
                                    "y0": s["bbox"][1],
                                    "x1": s["bbox"][2],
                                    "y1": s["bbox"][3],
                                })

            if raw_spans:
                # Group by Y coordinate
                raw_spans.sort(key=lambda s: (s["y0"], s["x0"]))
                lines = []
                for s in raw_spans:
                    placed = False
                    for line in lines:
                        avg_y = sum(item["y0"] for item in line) / len(line)
                        if abs(s["y0"] - avg_y) <= 4.0:
                            line.append(s)
                            line.sort(key=lambda item: item["x0"])
                            placed = True
                            break
                    if not placed:
                        lines.append([s])

                # Check if multiple lines have multi-column alignment
                multi_col_lines = [l for l in lines if len(l) >= 2]
                if len(multi_col_lines) >= 2:
                    table_count += 1
                    sheet_title = f"Page_{page_idx + 1}"
                    ws = wb.create_sheet(title=sheet_title[:31])

                    for r_idx, line in enumerate(lines):
                        for c_idx, span in enumerate(line):
                            val_str = span["text"].strip()
                            cell = ws.cell(row=r_idx + 1, column=c_idx + 1, value=val_str)
                            cell.border = cell_border
                            if r_idx == 0:
                                cell.font = header_font
                                cell.fill = header_fill
                            else:
                                cell.font = cell_font

                    for col in ws.columns:
                        col_letter = get_column_letter(col[0].column)
                        max_len = max(len(str(c.value or "")) for c in col)
                        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)

    # If no tables detected, create single sheet with full structured text
    if table_count == 0:
        ws = wb.create_sheet(title="Extracted Data")
        row_cursor = 1
        for page_idx, page in enumerate(pdf_doc):
            ws.cell(row=row_cursor, column=1, value=f"--- PAGE {page_idx + 1} ---").font = header_font
            ws.cell(row=row_cursor, column=1).fill = header_fill
            row_cursor += 1
            for line in page.get_text().splitlines():
                if line.strip():
                    ws.cell(row=row_cursor, column=1, value=line.strip()).font = cell_font
                    row_cursor += 1
            row_cursor += 1
        ws.column_dimensions["A"].width = 60

    pdf_doc.close()
    wb.save(excel_path)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python pdf_to_excel.py <input_pdf> <output_excel>", file=sys.stderr)
        sys.exit(1)
    try:
        convert_pdf_to_excel(sys.argv[1], sys.argv[2])
        if os.path.exists(sys.argv[2]) and os.path.getsize(sys.argv[2]) > 100:
            print("SUCCESS: PDF to Excel conversion complete.")
            sys.exit(0)
        else:
            print("ERROR: Generated Excel file is empty.", file=sys.stderr)
            sys.exit(1)
    except Exception as e:
        print(f"ERROR: {str(e)}", file=sys.stderr)
        sys.exit(1)
